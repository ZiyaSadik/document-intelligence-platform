"""Build docs/solution_presentation.pptx.

Kept as a script for the same reason as the architecture diagram: the deck
describes the system, so it should be regenerable when the system changes.

    python scripts/generate_presentation.py

WARNING - the committed deck has diverged from this script. It was re-themed
after generation (green accent, Georgia titles) and then edited in place: the
dashboard and result screenshots were refreshed from the live service and the
closing slide gained the deployment URLs. Running this script will overwrite
all of that with the older blue layout. Regenerate only if you intend to
rebuild the deck from scratch, and re-apply the theme afterwards.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "solution_presentation.pptx"
ARCHITECTURE = ROOT / "docs" / "architecture.png"
SHOTS = ROOT / "docs" / "screenshots"

# Palette: the product's own accent blue, a deep navy for the dark slides, and
# the same green/amber the dashboard uses for PASS and warning states, so the
# deck and the running application read as one thing.
NAVY = RGBColor(0x0F, 0x1B, 0x3D)
NAVY_MID = RGBColor(0x1B, 0x2C, 0x5C)
BLUE = RGBColor(0x1F, 0x4F, 0xD8)
BLUE_SOFT = RGBColor(0xE8, 0xEE, 0xFC)
INK = RGBColor(0x1B, 0x20, 0x29)
MUTED = RGBColor(0x62, 0x6B, 0x7A)
MUTED_LIGHT = RGBColor(0xA9, 0xB4, 0xC6)
BORDER = RGBColor(0xE2, 0xE5, 0xEA)
GREEN = RGBColor(0x17, 0x69, 0x3A)
GREEN_SOFT = RGBColor(0xE7, 0xF6, 0xEC)
AMBER = RGBColor(0x8A, 0x5A, 0x00)
AMBER_SOFT = RGBColor(0xFF, 0xF6, 0xE2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CANVAS = RGBColor(0xF7, 0xF8, 0xFA)

TITLE_FONT = "Cambria"
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"

W, H = 13.333, 7.5


def deck() -> Presentation:
    presentation = Presentation()
    presentation.slide_width = Inches(W)
    presentation.slide_height = Inches(H)
    return presentation


def blank(presentation: Presentation, *, dark: bool = False):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(H)
    )
    background.fill.solid()
    background.fill.fore_color.rgb = NAVY if dark else CANVAS
    background.line.fill.background()
    background.shadow.inherit = False
    return slide


def box(
    slide, x, y, w, h, *, fill=WHITE, line=BORDER, radius=0.035, rounded=True
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    if rounded:
        shape.adjustments[0] = radius
    return shape


def text(
    slide, x, y, w, h, paragraphs, *, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT
):
    """paragraphs: list of dicts with text/size/bold/color/font/space_after."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0

    for index, spec in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = spec.get("align", align)
        if spec.get("space_before"):
            paragraph.space_before = Pt(spec["space_before"])
        paragraph.space_after = Pt(spec.get("space_after", 0))
        if spec.get("line_spacing"):
            paragraph.line_spacing = spec["line_spacing"]
        run = paragraph.add_run()
        run.text = spec["text"]
        font = run.font
        font.size = Pt(spec.get("size", 14))
        font.bold = spec.get("bold", False)
        font.italic = spec.get("italic", False)
        font.name = spec.get("font", BODY_FONT)
        font.color.rgb = spec.get("color", INK)
    return shape


def slide_title(slide, title, subtitle=None, *, dark=False):
    text(
        slide, 0.75, 0.5, W - 1.5, 0.7,
        [{
            "text": title, "size": 34, "bold": True, "font": TITLE_FONT,
            "color": WHITE if dark else NAVY,
        }],
    )
    if subtitle:
        text(
            slide, 0.75, 1.12, W - 1.5, 0.42,
            [{
                "text": subtitle, "size": 14,
                "color": MUTED_LIGHT if dark else MUTED,
            }],
        )


def badge(slide, x, y, label, *, fill=BLUE, fg=WHITE, size=0.42):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = label
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.name = BODY_FONT
    run.font.color.rgb = fg
    return shape


def bullets(lines, *, size=12.5, color=INK, gap=7, bold_lead=False):
    out = []
    for line in lines:
        out.append({
            "text": "•  " + line, "size": size, "color": color,
            "space_after": gap, "bold": bold_lead,
        })
    return out


def footer(slide, note):
    text(
        slide, 0.75, H - 0.58, W - 1.5, 0.3,
        [{"text": note, "size": 10, "color": MUTED, "italic": True}],
    )


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------
def s_title(p):
    slide = blank(p, dark=True)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(10.4), Inches(-1.6), Inches(5.2), Inches(5.2)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = NAVY_MID
    accent.line.fill.background()
    accent.shadow.inherit = False

    text(
        slide, 1.0, 2.15, 9.6, 2.4,
        [
            {"text": "AI ENGINEER INTERNSHIP  ·  TECHNICAL CASE STUDY",
             "size": 12.5, "bold": True, "color": RGBColor(0x7E, 0x9C, 0xF0),
             "space_after": 16},
            {"text": "Document Intelligence Platform", "size": 44, "bold": True,
             "font": TITLE_FONT, "color": WHITE, "space_after": 10},
            {"text": "Intelligent extraction, financial validation and "
                     "API delivery for invoices and financial statements",
             "size": 16, "color": MUTED_LIGHT},
        ],
    )
    for index, (value, label) in enumerate(
        [("4", "document types"), ("140", "automated tests"),
         ("37", "checks, 0 failures"), ("1", "deployed service")]
    ):
        x = 1.0 + index * 2.45
        text(
            slide, x, 5.35, 2.2, 1.0,
            [
                {"text": value, "size": 30, "bold": True, "font": TITLE_FONT,
                 "color": WHITE, "space_after": 2},
                {"text": label, "size": 11, "color": MUTED_LIGHT},
            ],
        )


def s_brief(p):
    slide = blank(p)
    slide_title(
        slide, "The brief, and what was built",
        "One deployed service that turns a document into structured, "
        "reconciled, retrievable data",
    )

    box(slide, 0.75, 1.85, 5.9, 4.7)
    text(
        slide, 1.1, 2.15, 5.2, 4.1,
        [{"text": "WHAT WAS ASKED", "size": 11, "bold": True, "color": MUTED,
          "space_after": 14}]
        + bullets([
            "Accept PDF / JPG / PNG, up to 3 pages, and reject anything else "
            "gracefully",
            "Extract every visible field and table value — not a short "
            "fixed list — with evidence",
            "Validate the financial relationships for each document type, "
            "per comparative period",
            "Store results and serve them through REST APIs and a dashboard",
            "Deploy it: a localhost-only submission is incomplete",
        ]),
    )

    box(slide, 6.95, 1.85, 5.6, 4.7, fill=BLUE_SOFT, line=BLUE)
    text(
        slide, 7.3, 2.15, 4.9, 4.1,
        [{"text": "WHAT WAS DELIVERED", "size": 11, "bold": True, "color": BLUE,
          "space_after": 14}]
        + bullets([
            "FastAPI service with four documented endpoints and Swagger UI",
            "Text-layer-or-rasterise reader, then Claude vision with a "
            "schema-constrained response",
            "Deterministic reconciliation engine — no model in the "
            "arithmetic",
            "PostgreSQL-backed dashboard with raw-JSON view",
            "Docker image and Render blueprint; no system packages required",
        ]),
    )
    footer(slide, "Frontend and API are one process: one URL, one origin, no CORS.")


def s_finding(p):
    slide = blank(p)
    slide_title(
        slide, "The finding that shaped the design",
        "Investigating the corpus before writing code changed the architecture",
    )

    box(slide, 0.75, 1.9, 4.0, 3.1, fill=NAVY, line=None)
    text(
        slide, 1.1, 2.35, 3.3, 2.3,
        [
            {"text": "29 / 30", "size": 52, "bold": True, "font": TITLE_FONT,
             "color": WHITE, "space_after": 8},
            {"text": "statement PDFs have no text layer "
                     "and no embedded image",
             "size": 14, "color": MUTED_LIGHT},
        ],
    )

    box(slide, 5.05, 1.9, 7.5, 3.1)
    text(
        slide, 5.4, 2.2, 6.8, 2.6,
        [
            {"text": "Their text is painted as vector path outlines.", "size": 15,
             "bold": True, "space_after": 10},
            {"text": "Empty resource dictionary, no /Font, no /XObject — "
                     "roughly a megabyte of path-fill operators per page. "
                     "pdfplumber and pypdf both return zero characters, on "
                     "pages that look perfectly normal on screen.",
             "size": 13, "color": MUTED, "space_after": 12},
            {"text": "A text-layer-only solution extracts nothing from 29 of "
                     "the 30 files.",
             "size": 13.5, "bold": True, "color": AMBER},
        ],
    )

    consequences = [
        ("Rasterisation is the primary path",
         "not a fallback bolted on for scanned documents"),
        ("pypdfium2, not poppler",
         "a self-contained wheel, so the container needs no system packages"),
        ("Vision model over classical OCR",
         "row-to-column association in dense comparative tables is exactly "
         "where Tesseract-class OCR fails"),
    ]
    for index, (head, detail) in enumerate(consequences):
        x = 0.75 + index * 4.02
        box(slide, x, 5.2, 3.78, 1.5, fill=BLUE_SOFT, line=BLUE)
        text(
            slide, x + 0.3, 5.45, 3.2, 1.1,
            [
                {"text": head, "size": 12.5, "bold": True, "color": BLUE,
                 "space_after": 5},
                {"text": detail, "size": 10.5, "color": MUTED},
            ],
        )


def s_architecture(p):
    slide = blank(p)
    slide_title(
        slide, "Architecture",
        "Clear separation of concerns: each stage is one service with one job",
    )
    # The diagram is roughly 1.27:1 and the content area is 2.1:1, so it is
    # fitted by height; the freed width carries the points worth saying aloud.
    slide.shapes.add_picture(
        str(ARCHITECTURE), Inches(0.75), Inches(1.7), height=Inches(5.4)
    )

    notes = [
        ("One process, two surfaces",
         "The dashboard and the API are served by the same application, so "
         "there is one deploy, one URL and no CORS."),
        ("Each stage is swappable",
         "The reader, the extractor and the reconciler share no state; each "
         "is constructor-injected and separately testable."),
        ("One external call",
         "Only the extraction stage leaves the process. Everything else — "
         "validation, arithmetic, scoring — is local and deterministic."),
    ]
    for index, (head, detail) in enumerate(notes):
        y = 1.85 + index * 1.75
        box(slide, 7.85, y, 4.7, 1.55, fill=BLUE_SOFT, line=BLUE)
        text(
            slide, 8.15, y + 0.22, 4.1, 1.15,
            [
                {"text": head, "size": 12.5, "bold": True, "color": BLUE,
                 "space_after": 6},
                {"text": detail, "size": 10.5, "color": MUTED},
            ],
        )


def s_pipeline(p):
    slide = blank(p)
    slide_title(
        slide, "The pipeline, in five stages",
        "DocumentService owns the ordering and the failure handling; "
        "nothing else",
    )

    stages = [
        ("1", "Validate", "document_validation_service",
         ["Type from magic bytes, never the extension",
          "Empty, oversize, corrupted, > 3 pages",
          "Typed error → 400 / 413 / 422"], BLUE, BLUE_SOFT),
        ("2", "Read", "ocr_service",
         ["Native text layer when the page has one",
          "Otherwise rasterise at 200 DPI",
          "EXIF rotation, downscale, JPEG encode"], BLUE, BLUE_SOFT),
        ("3", "Extract", "extraction_service",
         ["One streamed Claude call, all pages",
          "JSON schema from Pydantic models",
          "Retry feeds the validator error back"], BLUE, BLUE_SOFT),
        ("4", "Reconcile", "financial_validation_service",
         ["Pure Python — no model involved",
          "Per document type, per period",
          "Missing operand → NOT_APPLICABLE"], GREEN, GREEN_SOFT),
        ("5", "Score", "confidence_service",
         ["Computed after the fact, not asked for",
          "Evidence corroboration is the key signal",
          "Absent value scores null, not zero"], AMBER, AMBER_SOFT),
    ]

    card_w, gap = 2.32, 0.18
    for index, (num, name, module, points, accent, soft) in enumerate(stages):
        x = 0.75 + index * (card_w + gap)
        box(slide, x, 1.95, card_w, 3.55, fill=soft, line=accent)
        badge(slide, x + 0.28, 2.22, num, fill=accent)
        text(
            slide, x + 0.28, 2.85, card_w - 0.56, 3.2,
            [
                {"text": name, "size": 17, "bold": True, "font": TITLE_FONT,
                 "color": accent, "space_after": 4},
                {"text": module, "size": 7, "font": MONO_FONT,
                 "color": MUTED, "space_after": 14},
            ]
            + bullets(points, size=10.5, gap=9),
        )

    footer(
        slide,
        "A rejected upload is still stored with processing_status FAILED, so "
        "the dashboard shows what was attempted rather than dropping it.",
    )


def s_extraction(p):
    slide = blank(p)
    slide_title(
        slide, "Extraction: the model transcribes, the code interprets",
        "The single most important design decision in the system",
    )

    box(slide, 0.75, 1.9, 5.75, 3.35, fill=BLUE_SOFT, line=BLUE)
    text(
        slide, 1.1, 2.2, 5.05, 2.8,
        [{"text": "THE MODEL RETURNS", "size": 11, "bold": True, "color": BLUE,
          "space_after": 12}]
        + bullets([
            "Every figure as a string, exactly as printed",
            "\"(1,234.56)\"   \"1,23,456.78\"   \"–\"",
            "The caption, section and page for each row",
            "The line of text each value came from",
            "null for anything it cannot read — never a guess",
        ], size=12.5),
    )

    box(slide, 6.8, 1.9, 5.75, 3.35, fill=GREEN_SOFT, line=GREEN)
    text(
        slide, 7.15, 2.2, 5.05, 2.8,
        [{"text": "THE APPLICATION DECIDES", "size": 11, "bold": True,
          "color": GREEN, "space_after": 12}]
        + bullets([
            "Parentheses mean negative",
            "Indian and Western digit grouping",
            "Currency symbols, Unicode and trailing minus",
            "Dash / Nil / N-A mean no value, not zero",
            "Every reconciliation, to the last decimal",
        ], size=12.5),
    )

    box(slide, 0.75, 5.5, 11.8, 1.2, fill=WHITE)
    text(
        slide, 1.1, 5.72, 11.1, 0.85,
        [
            {"text": "Why split it this way", "size": 12.5, "bold": True,
             "space_after": 5},
            {"text": "Numeric interpretation stays deterministic and "
                     "unit-tested instead of depending on the model doing "
                     "arithmetic; the as-printed form is preserved for audit; "
                     "and the confidence model can verify that a figure really "
                     "appears in the evidence text the model quoted.",
             "size": 12, "color": MUTED},
        ],
    )


def s_validation(p):
    slide = blank(p)
    slide_title(
        slide, "Financial validation",
        "Deterministic, per comparative period, and never inventing an operand",
    )

    groups = [
        ("Invoice / receipt", BLUE, BLUE_SOFT, [
            "quantity × unit_price = line amount",
            "Σ line amounts = subtotal",
            "subtotal + tax − discount = total",
            "cash_paid − total = change",
        ]),
        ("Balance sheet", BLUE, BLUE_SOFT, [
            "total capital and liabilities = total assets",
            "each section's components sum to that",
            "section's own reported total",
            "— derived from the document's structure",
        ]),
        ("Profit and loss", GREEN, GREEN_SOFT, [
            "interest earned + other income = total income",
            "interest expended + opex + provisions",
            "     = total expenditure",
            "income − expenditure = profit before MI",
            "PBMI − minority interest = attributable",
        ]),
        ("Cash flow", GREEN, GREEN_SOFT, [
            "operating + investing + financing + FX",
            "     = net increase in cash",
            "opening + net increase + cash acquired",
            "     = closing cash",
        ]),
    ]
    for index, (name, accent, soft, rules) in enumerate(groups):
        x = 0.75 + (index % 2) * 6.05
        y = 1.9 + (index // 2) * 2.05
        box(slide, x, y, 5.75, 1.85, fill=soft, line=accent)
        text(
            slide, x + 0.3, y + 0.22, 5.15, 1.45,
            [{"text": name, "size": 13, "bold": True, "color": accent,
              "space_after": 8}]
            + [{"text": rule, "size": 10.5, "font": MONO_FONT, "color": INK,
                "space_after": 3} for rule in rules],
        )

    box(slide, 0.75, 6.05, 11.8, 0.95, fill=WHITE)
    text(
        slide, 1.1, 6.2, 11.1, 0.7,
        [
            {"text": "Tolerance   |calc − reported|  ≤  0.1  +  1e-5 × magnitude"
                     "   — a combined allowance, both terms configurable",
             "size": 12, "bold": True, "space_after": 4},
            {"text": "An OR of absolute-or-relative is too loose at both ends: a "
                     "0.5% relative limit passes a 20,000-crore discrepancy, while "
                     "an absolute limit sized for a statement is 11% of a 9.00 till "
                     "receipt. A missing operand returns NOT_APPLICABLE and never "
                     "fails the document.",
             "size": 10.5, "color": MUTED},
        ],
    )


def s_confidence(p):
    slide = blank(p)
    slide_title(
        slide, "Confidence that means something",
        "Optional in the brief — implemented, but never asked of the model",
    )

    box(slide, 0.75, 1.95, 7.4, 3.6)
    rows = [
        ("Source path", "base 0.75 / 0.60",
         "A PDF text layer is exact; a rendered page was read visually."),
        ("Evidence corroboration", "+0.20",
         "The number is actually present in the quoted source text."),
        ("Evidence completeness", "+0.05 each",
         "A page number and a verbatim quote were both returned."),
        ("Validation agreement", "+0.10 / −0.30",
         "A figure inside a failing reconciliation is suspect."),
    ]
    text(
        slide, 1.1, 2.2, 6.7, 0.3,
        [{"text": "SIGNAL                                              WEIGHT",
          "size": 10, "bold": True, "color": MUTED}],
    )
    for index, (signal, weight, why) in enumerate(rows):
        y = 2.62 + index * 0.72
        text(
            slide, 1.1, y, 3.0, 0.6,
            [{"text": signal, "size": 12, "bold": True, "space_after": 3},
             {"text": why, "size": 9.5, "color": MUTED}],
        )
        text(
            slide, 4.35, y, 1.5, 0.3,
            [{"text": weight, "size": 12, "font": MONO_FONT, "color": BLUE}],
        )

    box(slide, 8.45, 1.95, 4.1, 3.6, fill=AMBER_SOFT, line=AMBER)
    text(
        slide, 8.8, 2.25, 3.4, 3.0,
        [
            {"text": "Why not ask the model?", "size": 14, "bold": True,
             "font": TITLE_FONT, "color": AMBER, "space_after": 10},
            {"text": "A self-reported score is not evidence. It correlates "
                     "with fluency, not with correctness, and it cannot be "
                     "checked.",
             "size": 12, "color": INK, "space_after": 10},
            {"text": "Every signal here is independently verifiable from the "
                     "response itself — which is what makes the number "
                     "explainable.",
             "size": 12, "color": INK},
        ],
    )

    box(slide, 0.75, 5.75, 11.8, 0.95, fill=WHITE)
    text(
        slide, 1.1, 5.92, 11.1, 0.65,
        [
            {"text": "A field the document does not report scores null, "
                     "not zero.", "size": 12, "bold": True, "space_after": 4},
            {"text": "There is nothing to be confident about, and zero would "
                     "read as “certainly wrong”. Values below 0.70 "
                     "are highlighted in the dashboard.",
             "size": 10.5, "color": MUTED},
        ],
    )


def s_api(p):
    slide = blank(p)
    slide_title(
        slide, "API surface",
        "Swagger generated from the same Pydantic models the endpoints return",
    )

    endpoints = [
        ("POST", "/api/v1/documents/process", "multipart upload; file + document_type"),
        ("GET", "/api/v1/documents/{document_name}", "latest stored result, case-insensitive"),
        ("GET", "/api/v1/documents", "dashboard list; filter, search, paginate"),
        ("GET", "/api/v1/health", "status, database, extraction_configured"),
    ]
    box(slide, 0.75, 1.95, 7.4, 2.5)
    for index, (verb, path, note) in enumerate(endpoints):
        y = 2.2 + index * 0.55
        text(slide, 1.1, y, 0.7, 0.3,
             [{"text": verb, "size": 11, "bold": True, "color": BLUE}])
        text(slide, 1.85, y, 4.2, 0.3,
             [{"text": path, "size": 11, "font": MONO_FONT}])
        text(slide, 1.85, y + 0.22, 6.0, 0.25,
             [{"text": note, "size": 9.5, "color": MUTED}])

    box(slide, 8.45, 1.95, 4.1, 2.5, fill=NAVY, line=None)
    text(
        slide, 8.8, 2.2, 3.4, 2.0,
        [
            {"text": "One error envelope", "size": 13, "bold": True,
             "color": WHITE, "space_after": 8},
            {"text": '{ "error": {\n    "code": "...",\n    "message": "...",\n'
                     '    "request_id": "..." } }',
             "size": 9.5, "font": MONO_FONT, "color": MUTED_LIGHT,
             "space_after": 8},
            {"text": "14 typed codes. Stack traces and provider messages go to "
                     "the log, never the body — asserted by a test.",
             "size": 10, "color": MUTED_LIGHT},
        ],
    )

    box(slide, 0.75, 4.65, 11.8, 2.05)
    text(
        slide, 1.1, 4.9, 11.1, 1.6,
        [
            {"text": "THE STORED RESPONSE", "size": 11, "bold": True,
             "color": MUTED, "space_after": 10}]
        + bullets([
            "extracted_data — every field with its parsed value, "
            "as-printed form, page number and evidence quote; line items; "
            "per-period key figures; an explicit missing_fields list",
            "validation — per check: formula, operands, calculated, "
            "reported, variance, PASS / FAIL / NOT_APPLICABLE, and the "
            "tolerance applied",
            "processing_metadata — ocr_used, extraction method, model, "
            "pages, timestamp, elapsed milliseconds",
        ], size=11.5, gap=6),
    )
    footer(
        slide,
        "The full response is persisted verbatim, so GET replays exactly what "
        "POST returned.",
    )


def s_dashboard(p):
    slide = blank(p)
    slide_title(
        slide, "The dashboard",
        "HTML, hand-written CSS and vanilla JS, calling the same public API",
    )
    if (SHOTS / "dashboard.png").exists():
        slide.shapes.add_picture(
            str(SHOTS / "dashboard.png"), Inches(0.75), Inches(1.95),
            width=Inches(6.9),
        )
    if (SHOTS / "document_result.png").exists():
        slide.shapes.add_picture(
            str(SHOTS / "document_result.png"), Inches(7.95), Inches(1.95),
            height=Inches(4.35),
        )
    text(
        slide, 0.75, 6.45, 5.9, 0.5,
        [{"text": "Upload, then every processed document with status, "
                  "validation outcome and confidence.",
          "size": 10.5, "color": MUTED}],
    )
    text(
        slide, 7.95, 6.45, 4.6, 0.5,
        [{"text": "Per document: extracted fields, per-period tables, every "
                  "check with its arithmetic, and the raw JSON.",
          "size": 10.5, "color": MUTED}],
    )


def s_quality(p):
    slide = blank(p)
    slide_title(
        slide, "Engineering discipline",
        "The parts that do not show up in a screenshot",
    )

    cells = [
        ("Testing", "140 tests, zero network",
         "The Anthropic client is stubbed throughout, so the suite runs the "
         "same in CI and with no API key."),
        ("Logging", "Structured, correlated",
         "Single-line JSON with an x-request-id on every record, returned in "
         "the response header and in error bodies."),
        ("Error handling", "One envelope, 14 codes",
         "Typed exceptions map to HTTP status; a test asserts no internal "
         "detail reaches the client."),
        ("Security", "Bytes over claims",
         "File type from magic bytes, filename sanitised against traversal, "
         "size ceiling before read, container runs unprivileged."),
        ("Configuration", "Environment only",
         "No secret has a default and none is committed; .env.example lists "
         "names, never values."),
        ("Modularity", "One job per module",
         "Validation, reading, extraction, reconciliation, scoring, "
         "persistence and routing are separate and separately testable."),
    ]
    for index, (name, headline, detail) in enumerate(cells):
        x = 0.75 + (index % 3) * 4.03
        y = 1.95 + (index // 3) * 2.35
        box(slide, x, y, 3.78, 2.1)
        badge(slide, x + 0.3, y + 0.26, "●", fill=BLUE_SOFT, fg=BLUE,
              size=0.34)
        text(
            slide, x + 0.78, y + 0.28, 2.7, 0.3,
            [{"text": name, "size": 13, "bold": True, "font": TITLE_FONT,
              "color": NAVY}],
        )
        text(
            slide, x + 0.3, y + 0.78, 3.2, 1.2,
            [
                {"text": headline, "size": 11.5, "bold": True, "color": BLUE,
                 "space_after": 6},
                {"text": detail, "size": 10.5, "color": MUTED},
            ],
        )


def s_deploy(p):
    slide = blank(p)
    slide_title(
        slide, "Deployment",
        "One Docker image, no system packages, one service serving both "
        "frontend and API",
    )

    box(slide, 0.75, 1.95, 5.75, 2.5, fill=BLUE_SOFT, line=BLUE)
    text(
        slide, 1.1, 2.2, 5.05, 2.0,
        [{"text": "WHY THE IMAGE IS SIMPLE", "size": 11, "bold": True,
          "color": BLUE, "space_after": 12}]
        + bullets([
            "Every dependency ships as a binary wheel",
            "No apt-get layer, no poppler, no compiler",
            "python:3.12-slim, runs as an unprivileged user",
            "render.yaml provisions service + database",
        ], size=12),
    )

    box(slide, 6.8, 1.95, 5.75, 2.5, fill=AMBER_SOFT, line=AMBER)
    text(
        slide, 7.15, 2.2, 5.05, 2.0,
        [
            {"text": "THE TRAP WORTH NAMING", "size": 11, "bold": True,
             "color": AMBER, "space_after": 12},
            {"text": "Free-tier container disks are ephemeral.", "size": 13,
             "bold": True, "space_after": 8},
            {"text": "A deployed SQLite file is wiped on every redeploy and on "
                     "cold start, so the dashboard would quietly empty itself "
                     "part-way through an evaluation. The ORM layer is "
                     "backend-agnostic; only DATABASE_URL changes, and the "
                     "blueprint points it at managed PostgreSQL.",
             "size": 11, "color": MUTED},
        ],
    )

    box(slide, 0.75, 4.7, 11.8, 2.0)
    text(
        slide, 1.1, 4.95, 11.1, 1.6,
        [{"text": "VERIFIED AFTER DEPLOY", "size": 11, "bold": True,
          "color": MUTED, "space_after": 10}]
        + bullets([
            "/api/v1/health returns ok, database connected, "
            "extraction_configured true",
            "Swagger UI loads; POST processes a real document; GET by name "
            "returns it; the list backs the dashboard",
            "An unsupported file returns 400 with the documented error "
            "envelope and no stack trace",
        ], size=11.5, gap=6),
    )


def s_limits(p):
    slide = blank(p)
    slide_title(
        slide, "Limitations, and what production would need",
        "Stated plainly — knowing where the edges are is part of the work",
    )

    box(slide, 0.75, 1.95, 5.75, 4.4, fill=AMBER_SOFT, line=AMBER)
    text(
        slide, 1.1, 2.25, 5.05, 3.9,
        [{"text": "KNOWN LIMITATIONS", "size": 11, "bold": True, "color": AMBER,
          "space_after": 14}]
        + bullets([
            "Accuracy is bounded by the vision model. The reconciliations "
            "catch internally inconsistent extractions; they cannot catch a "
            "value that is wrong self-consistently.",
            "The caption registry is English and oriented to the statement "
            "conventions in scope; an unusual caption leaves a check "
            "NOT_APPLICABLE rather than guessing.",
            "Processing is synchronous — 15 to 40 seconds per rasterised "
            "page. Fine for this brief, not for real concurrency.",
            "Confidence is an explainable heuristic, not calibrated against "
            "labelled ground truth.",
            "No authentication; the evaluation flow needs an open API.",
        ], size=11, gap=8),
    )

    box(slide, 6.8, 1.95, 5.75, 4.4, fill=GREEN_SOFT, line=GREEN)
    text(
        slide, 7.15, 2.25, 5.05, 3.9,
        [{"text": "WHAT I WOULD CHANGE FOR PRODUCTION", "size": 11, "bold": True,
          "color": GREEN, "space_after": 14}]
        + bullets([
            "Asynchronous processing: 202 plus a job id, a queue, and "
            "streamed progress.",
            "Content-hash caching so an identical re-upload costs nothing.",
            "Store the source document and link evidence to bounding boxes, so "
            "the UI highlights the exact region.",
            "An accuracy eval harness in CI — labelled corpus, field-level "
            "precision and recall — so a prompt change cannot silently "
            "regress extraction.",
            "Auth, per-tenant isolation, rate limiting and a model spend cap.",
            "Alembic migrations, OpenTelemetry traces, per-request cost "
            "accounting.",
        ], size=11, gap=8),
    )


def s_close(p):
    slide = blank(p, dark=True)
    text(
        slide, 1.0, 1.6, 11.3, 1.2,
        [
            {"text": "Thank you", "size": 40, "bold": True, "font": TITLE_FONT,
             "color": WHITE, "space_after": 10},
            {"text": "Happy to walk through the code, the prompts, the "
                     "validation logic or the deployment — and to change "
                     "something live.",
             "size": 15, "color": MUTED_LIGHT},
        ],
    )

    links = [
        ("Frontend", "the deployed dashboard"),
        ("Backend API", "/api/v1"),
        ("Swagger / OpenAPI", "/docs"),
        ("Health", "/api/v1/health"),
        ("Repository", "public GitHub"),
    ]
    for index, (label, note) in enumerate(links):
        y = 3.5 + index * 0.62
        text(
            slide, 1.0, y, 3.2, 0.4,
            [{"text": label, "size": 13, "bold": True, "color": WHITE}],
        )
        text(
            slide, 4.4, y, 7.5, 0.4,
            [{"text": note, "size": 12, "color": MUTED_LIGHT,
              "font": MONO_FONT}],
        )

    box(slide, 8.6, 3.4, 3.7, 3.0, fill=NAVY_MID, line=None)
    text(
        slide, 8.95, 3.7, 3.0, 2.5,
        [
            {"text": "Also in the repository", "size": 12, "bold": True,
             "color": WHITE, "space_after": 10},
            {"text": "README with setup, API reference, validation rules and "
                     "tolerance\n\nArchitecture diagram\n\nDeployment "
                     "runbook\n\nSample JSON outputs from the real pipeline\n\n"
                     "140 automated tests",
             "size": 10.5, "color": MUTED_LIGHT},
        ],
    )


def main() -> int:
    if not ARCHITECTURE.exists():
        print(
            "docs/architecture.png is missing; run "
            "scripts/generate_architecture.py first",
            file=sys.stderr,
        )
        return 2

    presentation = deck()
    for builder in (
        s_title, s_brief, s_finding, s_architecture, s_pipeline, s_extraction,
        s_validation, s_confidence, s_api, s_dashboard, s_quality, s_deploy,
        s_limits, s_close,
    ):
        builder(presentation)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    print(
        f"wrote {OUTPUT.relative_to(ROOT)} "
        f"({len(presentation.slides.__iter__.__self__._sldIdLst)} slides, "
        f"{OUTPUT.stat().st_size // 1024} KB)"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
